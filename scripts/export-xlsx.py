#!/usr/bin/env python3
"""Baut daten/Ausbildungsbetriebe.xlsx aus daten/betriebe.json.

Wird vom täglichen Sammellauf (GitHub Actions) automatisch mit ausgeführt,
damit im Repository immer eine fertig formatierte Excel-Übersicht liegt und
niemand die rohe CSV auf GitHub selbst zusammenklicken muss.

Lokal von Hand aufrufen, z. B. nach einem manuellen Import:
    pip install openpyxl
    python3 scripts/export-xlsx.py
"""
import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJEKT = Path(__file__).resolve().parent.parent
QUELLE = PROJEKT / "daten" / "betriebe.json"
ZIEL = PROJEKT / "daten" / "Ausbildungsbetriebe.xlsx"

RUHEND_AB_TAGEN = 60
SCHRIFT = "Arial"


def tage_seit(datum_str):
    if not datum_str:
        return None
    jahr, monat, tag = map(int, datum_str.split("-"))
    return (date.today() - date(jahr, monat, tag)).days


def zeile_aus_betrieb(betrieb):
    berufe = betrieb.get("berufe", [])
    tage = tage_seit(betrieb.get("zuletztGesehen"))
    return {
        "Arbeitgeber": betrieb.get("arbeitgeber", ""),
        "Ort": betrieb.get("ort", ""),
        "PLZ": betrieb.get("plz", ""),
        "Strasse": betrieb.get("strasse", ""),
        "Region": betrieb.get("region", ""),
        "Ausbildungsberufe": " | ".join(e.get("beruf", "") for e in berufe),
        "AnzahlBerufe": len(berufe),
        "AnzahlAnzeigen": betrieb.get("anzahlAnzeigen", 0),
        "ErstmalsGesehen": betrieb.get("erstmalsGesehen", ""),
        "ZuletztGesehen": betrieb.get("zuletztGesehen", ""),
        "TageOhneAnzeige": tage if tage is not None else "",
        "Status": (
            "aktuell ausgeschrieben"
            if tage is not None and tage < RUHEND_AB_TAGEN
            else "ohne aktuelle Anzeige"
        ),
        "Jahre": " ".join(betrieb.get("jahre", [])),
        "Quellen": " ".join(betrieb.get("quellen", [])),
        "Notiz": betrieb.get("notiz", ""),
    }


def main():
    bestand = json.loads(QUELLE.read_text(encoding="utf-8"))
    betriebe = bestand["betriebe"]

    zeilen = sorted(
        (zeile_aus_betrieb(b) for b in betriebe),
        key=lambda z: (z["Ort"], z["Arbeitgeber"]),
    )
    spalten = list(zeilen[0].keys()) if zeilen else list(zeile_aus_betrieb({}).keys())

    wb = Workbook()
    ws = wb.active
    ws.title = "Ausbildungsbetriebe"

    kopf_font = Font(name=SCHRIFT, bold=True, color="FFFFFF", size=11)
    kopf_fill = PatternFill("solid", fgColor="8A5A2B")
    normal_font = Font(name=SCHRIFT, size=10)
    ruhend_fill = PatternFill("solid", fgColor="FCEFD9")

    ws.append(spalten)
    for zelle in ws[1]:
        zelle.font = kopf_font
        zelle.fill = kopf_fill
        zelle.alignment = Alignment(vertical="center")

    status_index = spalten.index("Status")
    for zeile in zeilen:
        ws.append([zeile[s] for s in spalten])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        ruhend = row[status_index].value == "ohne aktuelle Anzeige"
        for zelle in row:
            zelle.font = normal_font
            if ruhend:
                zelle.fill = ruhend_fill

    breiten = {
        "Arbeitgeber": 34, "Ort": 16, "PLZ": 7, "Strasse": 22, "Region": 20,
        "Ausbildungsberufe": 46, "AnzahlBerufe": 12, "AnzahlAnzeigen": 13,
        "ErstmalsGesehen": 14, "ZuletztGesehen": 14, "TageOhneAnzeige": 14,
        "Status": 20, "Jahre": 12, "Quellen": 14, "Notiz": 28,
    }
    for i, name in enumerate(spalten, start=1):
        ws.column_dimensions[get_column_letter(i)].width = breiten.get(name, 14)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spalten))}{ws.max_row}"
    ws.sheet_view.showGridLines = False

    info = wb.create_sheet("Info")
    info["A1"] = "Ausbildungsbetriebe der Region – automatisch gesammelt aus der BA-Jobsuche"
    info["A1"].font = Font(name=SCHRIFT, bold=True, size=12)
    hinweise = [
        "",
        f"Stand: {bestand.get('aktualisiert', '')}",
        f"Anzahl Betriebe: {len(betriebe)}",
        "",
        "Ein Betrieb bleibt in der Liste, auch wenn er gerade nichts ausgeschrieben hat –",
        '„Status = ohne aktuelle Anzeige" ist für Initiativbewerbungen die interessante Gruppe.',
        "(gelb hinterlegt in der Tabelle)",
        "",
        'Quelle „ba-jobsuche": aus der Jobbörse der Bundesagentur für Arbeit.',
        "Weitere Quellen (z. B. hwk-stuttgart) kommen über Kammer-Importe dazu.",
        "",
        "Diese Datei wird täglich automatisch neu erzeugt. Aktuelle Version immer im Repository:",
        "https://github.com/Penetrationpasi-hub/Clerking",
    ]
    for i, zeile in enumerate(hinweise, start=2):
        info[f"A{i}"] = zeile
        info[f"A{i}"].font = Font(name=SCHRIFT, size=10)
    info.column_dimensions["A"].width = 85

    wb.save(ZIEL)
    print(f"Geschrieben: {ZIEL} ({len(zeilen)} Betriebe)")


if __name__ == "__main__":
    main()
